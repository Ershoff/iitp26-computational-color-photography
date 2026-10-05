# -*- coding: utf-8 -*-
"""Собрать черновик Lec_3 в дизайне Lec_1,2 (Simple Light, Lato, 16:9 10\").

Taught deck: Презентации/Lec_3_linear_model.pptx
(правка 30.09.2026: титул ТМШ, 55 слайдов). Этот скрипт пишет
Lec_3_linear_model_built.pptx и не затирает taught-файл.
"""
from __future__ import annotations

from pathlib import Path
from shutil import copyfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

plt.rcParams["font.family"] = "Arial"
plt.rcParams["axes.unicode_minus"] = False

ROOT = Path(__file__).resolve().parents[2]
FIG = Path(__file__).resolve().parent / "figures"
PHOTO = FIG / "photo"
CROP = FIG / "crop"
KHARK = FIG / "kharkevich"
L2FIG = ROOT / "weeks" / "02" / "figures"
SRC = ROOT / "Презентации" / "Lec_1,2_intro2color.pptx"
OUT = ROOT / "Презентации" / "Lec_3_linear_model_built.pptx"

BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)
BLUE = RGBColor(0x42, 0x85, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xEE, 0xEE, 0xEE)
FONT = "Lato"

TITLE_BOX = (0.355, 0.463, 9.183, 0.695)
BODY_BOX = (0.912, 1.236, 8.112, 3.549)
BODY_LEFT_W = 4.15
PIC_RIGHT = (5.273, 1.269, 4.375, 3.306)
QUESTION_BOX = (0.912, 2.153, 8.112, 1.481)
CAPTION_BOX = (0.447, 5.188, 9.1, 0.387)

P_CXNSP = "{http://schemas.openxmlformats.org/presentationml/2006/main}cxnSp"

L12_FILES = {
    "s15_Google_Shape_247_p47.png": "l12_brdf.png",
    "s15_Google_Shape_249_p47.png": "l12_eye.png",
    "s15_Google_Shape_250_p47.png": "l12_lms.png",
    "s16_Google_Shape_257_p48.png": "l12_lettuce.png",
    "s16_Google_Shape_259_p48.png": "l12_lettuce_ink.png",
    "s24_Google_Shape_328_p56.png": "l12_match.png",
    "s25_Google_Shape_336_p57.png": "l12_cmf.png",
    "s27_Google_Shape_353_p59.png": "l12_cone.png",
    "s27_Google_Shape_354_p59.png": "l12_xy.png",
    "s43_Google_Shape_471_p75.png": "l12_planck.png",
    "s43_Google_Shape_470_p75.png": "l12_planck_xy.png",
    "s44_Google_Shape_481_p76.png": "l12_cct.png",
    "s45_Google_Shape_488_p77.png": "l12_spectra.png",
    "s05_Google_Shape_171_p37.png": "l12_shoot.png",
}


def delete_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    ns = qn("r:id")
    for sld_id in list(sld_id_lst):
        r_id = sld_id.get(ns)
        prs.part.drop_rel(r_id)
        sld_id_lst.remove(sld_id)


def _layouts_named(prs: Presentation, name: str):
    found = []
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                found.append(layout)
    return found


def layout_named(prs: Presentation, name: str):
    found = _layouts_named(prs, name)
    if not found:
        raise KeyError(name)
    if name == "TITLE_ONLY":
        ruled = [ly for ly in found if ly._element.findall(f".//{P_CXNSP}")]
        if ruled:
            return ruled[0]
    return found[0]


def _run(run, size=None, bold=False, color=BLACK, font=FONT, italic=False):
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    # Иначе латинские n, f в русской строке берутся из ea/cs плейсхолдера.
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def _textbox(slide, box, text, size, bold=False, italic=False, color=BLACK, align=PP_ALIGN.LEFT):
    l, t, w, h = box
    shp = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _run(r, size, bold=bold, italic=italic, color=color)
    return shp


def _clear_runs(p):
    for r in list(p.runs):
        el = r._r
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _set_bullet(p, on: bool) -> None:
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        tag = child.tag
        if tag.endswith("}buFont") or tag.endswith("}buChar") or tag.endswith("}buAutoNum") or tag.endswith(
            "}buNone"
        ) or tag.endswith("}buBlip"):
            pPr.remove(child)
    if on:
        return
    etree.SubElement(pPr, qn("a:buNone"))


def body_shape(slide):
    for shp in slide.placeholders:
        try:
            if shp.placeholder_format.type == PP_PLACEHOLDER.BODY:
                return shp
        except Exception:
            continue
    return None


def set_body(slide, items, *, width=None, height=None, left=None, top=None, size=18, align=PP_ALIGN.LEFT, bold=False, bullets=True):
    """Текст в штатный body-плейсхолдер TITLE_ONLY — как в Lec_1,2."""
    if isinstance(items, str):
        items = [items]
    ph = body_shape(slide)
    l = BODY_BOX[0] if left is None else left
    t = BODY_BOX[1] if top is None else top
    w = BODY_BOX[2] if width is None else width
    h = BODY_BOX[3] if height is None else height
    if ph is None:
        shp = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = shp.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(10)
            _set_bullet(p, bullets)
            r = p.add_run()
            r.text = item
            _run(r, size, bold=bold)
        return shp
    # Все четыре числа сразу: иначе python-pptx пишет xfrm с нулевой стороной.
    ph.left, ph.top = Inches(l), Inches(t)
    ph.width, ph.height = Inches(w), Inches(h)
    tf = ph.text_frame
    tf.word_wrap = True
    while len(tf.paragraphs) > 1:
        el = tf.paragraphs[-1]._p
        el.getparent().remove(el)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.level = 0
        p.space_after = Pt(10)
        _clear_runs(p)
        _set_bullet(p, bullets)
        r = p.add_run()
        r.text = item
        _run(r, size, bold=bold)
    return ph


def place_title(slide, text):
    box = slide.shapes.title
    box.left, box.top = Inches(TITLE_BOX[0]), Inches(TITLE_BOX[1])
    box.width, box.height = Inches(TITLE_BOX[2]), Inches(TITLE_BOX[3])
    tf = box.text_frame
    p = tf.paragraphs[0]
    _clear_runs(p)
    r = p.add_run()
    r.text = text
    _run(r, 28, bold=True)


def titled(prs, title):
    slide = prs.slides.add_slide(layout_named(prs, "TITLE_ONLY"))
    place_title(slide, title)
    return slide


def content(prs, title, items):
    slide = titled(prs, title)
    set_body(slide, items)
    return slide


def section(prs, title, photo=None, note=None):
    """Вводный слайд: заголовок + фото Харкевича на правую треть."""
    slide = titled(prs, title)
    hide_body(slide)
    if photo is not None and Path(photo).exists():
        cover_pic(slide, photo, (6.70, 1.16, 3.15, 3.85))
        if note:
            items = [note] if isinstance(note, str) else note
            set_body(slide, items, width=5.85, size=18, bullets=len(items) > 1)
        caption(slide, "Фотоархив А. А. Харкевича, ИППИ РАН")
    elif note:
        items = [note] if isinstance(note, str) else note
        set_body(
            slide,
            items,
            left=QUESTION_BOX[0],
            top=QUESTION_BOX[1],
            width=QUESTION_BOX[2],
            height=QUESTION_BOX[3],
            size=30,
            bold=True,
            align=PP_ALIGN.CENTER,
            bullets=False,
        )
    return slide


def four_pic(prs, title, items, caption_txt=None):
    """Четыре фото 2×2 с подписью под каждым."""
    slide = titled(prs, title)
    hide_body(slide)
    gap_x, gap_y = 0.14, 0.08
    left, top = 0.40, 1.16
    cell_w, cell_h = 4.50, 1.85
    cap_h = 0.36
    for i, (path, label) in enumerate(items):
        col, row = i % 2, i // 2
        x = left + col * (cell_w + gap_x)
        y = top + row * (cell_h + gap_y)
        cover_pic(slide, path, (x, y, cell_w, cell_h - cap_h))
        _textbox(slide, (x, y + cell_h - cap_h, cell_w, cap_h), label, 13, align=PP_ALIGN.CENTER)
    if caption_txt:
        caption(slide, caption_txt)
    return slide


def title_slide(prs, title, subtitle, foot):
    slide = prs.slides.add_slide(layout_named(prs, "TITLE"))
    slide.shapes.title.text = title
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            _run(r, 32, bold=True)
    sub = slide.placeholders[1]
    tf = sub.text_frame
    p = tf.paragraphs[0]
    _clear_runs(p)
    r = p.add_run()
    r.text = subtitle
    _run(r, 20)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = foot
    _run(r2, 16, color=GRAY)
    return slide


def pause(slide, text):
    _textbox(slide, (0.912, 4.90, 8.112, 0.42), text, 14, italic=True, color=BLUE)


def caption(slide, text):
    _textbox(slide, CAPTION_BOX, text, 11, color=GRAY)


def pic(slide, path, left, top, width=None, height=None):
    kw = {}
    if width:
        kw["width"] = Inches(width)
    if height:
        kw["height"] = Inches(height)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kw)


_COVER_N = 0


def cover_pic(slide, path, box):
    global _COVER_N
    l, t, w, h = box
    im = Image.open(path).convert("RGB")
    tw, th = max(2, int(w * 140)), max(2, int(h * 140))
    iw, ih = im.size
    scale = max(tw / iw, th / ih)
    nw, nh = max(tw, int(iw * scale)), max(th, int(ih * scale))
    im = im.resize((nw, nh), Image.Resampling.LANCZOS)
    x0, y0 = (nw - tw) // 2, (nh - th) // 2
    im = im.crop((x0, y0, x0 + tw, y0 + th))
    _COVER_N += 1
    tmp = FIG / f"_cover_{_COVER_N}.jpg"
    im.save(tmp, quality=88)
    slide.shapes.add_picture(str(tmp), Inches(l), Inches(t), Inches(w), Inches(h))


def fit_pic(slide, path, box):
    l, t, w, h = box
    im = Image.open(path)
    iw, ih = im.size
    aspect = iw / ih
    box_a = w / h
    if aspect > box_a:
        nw, nh = w, w / aspect
        x, y = l, t + (h - nh) / 2
    else:
        nh, nw = h, h * aspect
        y, x = t, l + (w - nw) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(nw), Inches(nh))


def text_pic(prs, title, items, image, caption_txt=None, pause_txt=None, cover=False):
    """Текст слева, фигура справа — как «Что думают биологи?»."""
    slide = titled(prs, title)
    set_body(slide, items, width=BODY_LEFT_W)
    box = PIC_RIGHT
    if cover:
        cover_pic(slide, image, box)
    else:
        fit_pic(slide, image, box)
    if caption_txt:
        caption(slide, caption_txt)
    if pause_txt:
        pause(slide, pause_txt)
    return slide


def hide_body(slide):
    ph = body_shape(slide)
    if ph is None:
        return
    ph.left, ph.top = Inches(0.01), Inches(5.50)
    ph.width, ph.height = Inches(0.10), Inches(0.10)


def pic_slide(prs, title, image, caption_txt=None, box=(0.47, 1.16, 9.10, 3.70)):
    slide = titled(prs, title)
    hide_body(slide)
    fit_pic(slide, image, box)
    if caption_txt:
        caption(slide, caption_txt)
    return slide


def two_pic(prs, title, p1, p2, caption_txt=None, b1=(0.47, 1.16, 5.20, 3.70), b2=(5.85, 1.22, 3.70, 3.55)):
    slide = titled(prs, title)
    hide_body(slide)
    fit_pic(slide, p1, b1)
    fit_pic(slide, p2, b2)
    if caption_txt:
        caption(slide, caption_txt)
    return slide


def formula_png(path: Path, tex: str, fontsize=24, h=1.15):
    fig, ax = plt.subplots(figsize=(8.2, h))
    ax.axis("off")
    ax.text(0.5, 0.5, tex, ha="center", va="center", fontsize=fontsize)
    fig.savefig(path, dpi=180, bbox_inches="tight", pad_inches=0.06, facecolor="white")
    plt.close(fig)


def formula_stack_png(path: Path, lines, fontsize=18, h=1.55):
    fig, ax = plt.subplots(figsize=(8.4, h))
    ax.axis("off")
    n = len(lines)
    if n == 1:
        ys = [0.5]
    else:
        ys = [1.0 - (i + 0.55) / (n + 0.1) for i in range(n)]
    for y, line in zip(ys, lines):
        ax.text(0.5, y, line, ha="center", va="center", fontsize=fontsize)
    fig.savefig(path, dpi=180, bbox_inches="tight", pad_inches=0.06, facecolor="white")
    plt.close(fig)


def _clean(ax):
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)


def make_figures():
    FIG.mkdir(exist_ok=True)
    CROP.mkdir(exist_ok=True)
    dump = FIG / "from_l12"
    for src_name, dest_name in L12_FILES.items():
        dest = FIG / dest_name
        src = dump / src_name
        if not dest.exists() and src.exists():
            copyfile(src, dest)
    for name in ("lms_cones.png", "maxwell_match.png", "lettuce_metamers.png"):
        src = L2FIG / name
        if src.exists() and not (FIG / name).exists():
            (FIG / name).write_bytes(src.read_bytes())

    formula_png(FIG / "eq_sensor.png", r"$c_i=\int F(\lambda)\,\chi_i(\lambda)\,d\lambda$")
    formula_stack_png(
        FIG / "eq_full_F.png",
        [
            r"$F(\mathbf{n}_r,\lambda)=\sum_j\sum_k\int S_j(\lambda)\,B_j(\mathbf{n}_i)$",
            r"$\qquad\times R_k(\mathbf{n}_i,\mathbf{n}_r,\lambda)\,f(\mathbf{n}_i,\mathbf{n}_r)\,d\omega_i$",
        ],
        fontsize=16,
        h=1.15,
    )
    formula_stack_png(
        FIG / "eq_full_stack.png",
        [
            r"$F(\mathbf{n}_r,\lambda)=\sum_j\sum_k\int S_j(\lambda)\,B_j(\mathbf{n}_i)\,R_k(\mathbf{n}_i,\mathbf{n}_r,\lambda)\,f(\mathbf{n}_i,\mathbf{n}_r)\,d\omega_i$",
            r"$c_i(\mathbf{n}_r)=\int F(\mathbf{n}_r,\lambda)\,\chi_i(\lambda)\,d\lambda$",
        ],
        fontsize=13,
        h=1.35,
    )
    formula_png(
        FIG / "eq_sensor_r.png",
        r"$c_i(\mathbf{n}_r)=\int F(\mathbf{n}_r,\lambda)\,\chi_i(\lambda)\,d\lambda$",
    )
    formula_png(FIG / "eq_brdf.png", r"$R(\mathbf{n}_i,\mathbf{n}_r,\lambda)$")
    formula_png(FIG / "eq_diffuse.png", r"$F(\lambda)=S(\lambda)\,\Phi(\lambda)$")
    formula_png(
        FIG / "eq_complement.png",
        r"$\Phi^*(\lambda)=1-\Phi(\lambda),\quad C(S\Phi)+C(S\Phi^*)=C(S)$",
    )
    formula_png(FIG / "eq_delta.png", r"$\int \Delta(\lambda)\,\chi_i(\lambda)\,d\lambda=0,\quad \Delta=F-\Psi$")
    formula_png(FIG / "eq_mix.png", r"$\Phi=\sum_k \alpha_k\Phi_k,\quad \alpha_k\geq 0,\ \sum\alpha_k=1$")
    formula_png(FIG / "eq_piecewise.png", r"$\Phi=\sum_k a_k\,1_{I_k},\quad 0\leq a_k\leq 1$")
    formula_stack_png(
        FIG / "eq_kang.png",
        [
            r"$A^{T}=(V^{T}\Phi)^{-1}(V^{T}\Phi_j)\,A_j^{T}$",
            r"$A^{T}=(A^{T}\Phi_j)\,A_j^{T}$",
            r"$A_j^{T}=(A^{T}\Phi_j)^{-1}A^{T}$",
        ],
        fontsize=18,
        h=1.75,
    )

    lam = np.linspace(400, 700, 600)

    def _bb(T):
        x = lam * 1e-9
        h, c, k = 6.626e-34, 2.998e8, 1.381e-23
        B = 1.0 / (x**5 * (np.exp(np.minimum(h * c / (x * k * T), 80.0)) - 1))
        return B / np.nanmax(B)

    def _gauss(mu, sig):
        return np.exp(-0.5 * ((lam - mu) / sig) ** 2)

    def _norm(y):
        y = np.asarray(y, dtype=float)
        m = np.nanmax(y)
        return y / m if m > 0 else y

    fig, axes = plt.subplots(2, 2, figsize=(8.6, 4.15), sharex=True, sharey=True)
    panels = [
        (axes[0, 0], _bb(2856), "A / накал, Планк"),
        (
            axes[0, 1],
            0.35 * _gauss(540, 70)
            + 0.95 * _gauss(436, 3)
            + 0.75 * _gauss(546, 3)
            + 0.55 * _gauss(578, 3),
            "FL: линии + люминофор",
        ),
        (axes[1, 0], 0.95 * _gauss(450, 10) + 0.7 * _gauss(580, 55), "LED-B: синий + люминофор"),
        (
            axes[1, 1],
            0.55 * _gauss(410, 12) + 0.85 * _gauss(530, 75) + 0.45 * _gauss(620, 55),
            "LED-V: широкополосный",
        ),
    ]
    for ax, y, title in panels:
        ax.plot(lam, _norm(y), color="#111", lw=1.6)
        ax.set_title(title, fontsize=11)
        ax.set_xlim(400, 700)
        ax.set_ylim(0, 1.15)
        ax.set_xticks([400, 500, 600, 700])
        _clean(ax)
    axes[1, 0].set_xlabel(r"$\lambda$, nm")
    axes[1, 1].set_xlabel(r"$\lambda$, nm")
    axes[0, 0].set_ylabel("отн. спектр")
    axes[1, 0].set_ylabel("отн. спектр")
    fig.tight_layout()
    fig.savefig(FIG / "typical_sources.png", dpi=180, facecolor="white")
    plt.close(fig)

    def _lms(shift_L=0.0, width_scale=1.0):
        L = _gauss(570 + shift_L, 38 * width_scale)
        M = _gauss(543, 36 * width_scale)
        S = _gauss(445, 20)
        return L, M, S

    def _plot_lms(ax, L, M, S, title):
        ax.plot(lam, L, color="#c0392b", lw=1.5)
        ax.plot(lam, M, color="#1e8449", lw=1.5)
        ax.plot(lam, S, color="#2471a3", lw=1.5)
        ax.set_title(title, fontsize=11)
        ax.set_xlim(400, 700)
        ax.set_ylim(0, 1.15)
        ax.set_xticks([400, 500, 600, 700])
        _clean(ax)

    L0, M0, S0 = _lms()
    lens = 1 / (1 + np.exp(-(lam - 430) / 18))
    mac = np.exp(-1.15 * _gauss(460, 18))
    fig, axes = plt.subplots(2, 2, figsize=(8.2, 4.35), sharex=True, sharey=True)
    _plot_lms(axes[0, 0], L0, M0, S0, "средний CIEPO06")
    _plot_lms(axes[0, 1], L0 * lens, M0 * lens, S0 * lens, "линза режет синее")
    _plot_lms(axes[1, 0], L0 * mac, M0 * mac, S0 * mac, "макула, яма ~460 нм")
    Lw, Mw, Sw = _lms(shift_L=12, width_scale=1.25)
    _plot_lms(axes[1, 1], Lw, Mw, Sw, "сдвиг λmax и плотность L")
    axes[1, 0].set_xlabel(r"$\lambda$, nm")
    axes[1, 1].set_xlabel(r"$\lambda$, nm")
    axes[0, 0].set_ylabel(r"отн. $\chi$")
    axes[1, 0].set_ylabel(r"отн. $\chi$")
    fig.tight_layout()
    fig.savefig(FIG / "asano_chi.png", dpi=180, facecolor="white")
    plt.close(fig)

    plt.rcParams.update({"font.size": 11, "axes.linewidth": 0.8})
    fig, axes = plt.subplots(1, 3, figsize=(9.4, 2.55), sharey=True)
    a = np.exp(-((lam - 540) ** 2) / 2 / 55**2)
    pairs = [
        (0.95 * a, 0.45 * a, "0 пересечений"),
        (0.9 * a, np.clip(0.55 * a + 0.35 * np.tanh((lam - 555) / 8), 0, 1), "1 пересечение"),
        (0.85 * a, np.clip(0.55 + 0.4 * np.sin((lam - 400) / 28), 0, 1), "3 пересечения"),
    ]
    for ax, (y1, y2, title) in zip(axes, pairs):
        ax.plot(lam, y1, color="#111", lw=1.6)
        ax.plot(lam, y2, color="#111", lw=1.6, ls="--")
        ax.set_title(title, fontsize=11)
        ax.set_xlim(400, 700)
        ax.set_ylim(0, 1.05)
        ax.set_xlabel(r"$\lambda$, nm")
        ax.set_xticks([400, 500, 600, 700])
        _clean(ax)
    axes[0].set_ylabel("отн. спектр")
    fig.tight_layout()
    fig.savefig(FIG / "crossings.png", dpi=180, facecolor="white")
    plt.close(fig)

    fig, axes = plt.subplots(1, 3, figsize=(9.4, 2.45), sharey=True)
    s0 = np.zeros_like(lam)
    s1 = (lam > 560).astype(float)
    s2 = ((lam > 490) & (lam < 610)).astype(float)
    for ax, y, title in zip(axes, (s0, s1, s2), (r"$m=0$", r"$m=1$", r"$m=2$")):
        ax.fill_between(lam, 0, y, color="#dddddd", step="mid")
        ax.step(lam, y, where="mid", color="#111", lw=1.7)
        ax.set_title(title, fontsize=11)
        ax.set_xlim(400, 700)
        ax.set_ylim(-0.05, 1.15)
        ax.set_xlabel(r"$\lambda$, nm")
        ax.set_xticks([400, 500, 600, 700])
        _clean(ax)
    axes[0].set_ylabel(r"$\Phi$")
    fig.tight_layout()
    fig.savefig(FIG / "steps.png", dpi=180, facecolor="white")
    plt.close(fig)


def slide_block_a_timeline(prs):
    """Горизонтальный таймлайн блока А: четыре лекции, сегодня — Л3."""
    slide = titled(prs, "Блок А")
    hide_body(slide)

    stops = [
        ("Л1", "Что такое цвет", "спектр, колбочки", "было", "past"),
        ("Л2", "Линейная модель", "Грассман, интеграл", "было", "past"),
        ("Л3", "Разбор формулы", "пять следствий", "сегодня", "now"),
        ("Л4", "Восприятие", "пороги, CAM", "дальше", "next"),
    ]
    y_line = 2.22
    r = 0.11
    x0, x1 = 1.15, 8.85
    n = len(stops)
    xs = [x0 + i * (x1 - x0) / (n - 1) for i in range(n)]

    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x0),
        Inches(y_line),
        Inches(x1),
        Inches(y_line),
    )
    conn.line.color.rgb = BLACK
    conn.line.width = Pt(1.5)

    col_w = 2.35
    for x, (num, title, sub, when, kind) in zip(xs, stops):
        oval = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x - r),
            Inches(y_line - r),
            Inches(2 * r),
            Inches(2 * r),
        )
        oval.line.color.rgb = BLUE if kind == "now" else BLACK
        oval.line.width = Pt(1.25)
        oval.fill.solid()
        if kind == "past":
            oval.fill.fore_color.rgb = BLACK
        elif kind == "now":
            oval.fill.fore_color.rgb = BLUE
        else:
            oval.fill.fore_color.rgb = WHITE

        when_color = BLUE if kind == "now" else GRAY
        _textbox(
            slide,
            (x - col_w / 2, 1.35, col_w, 0.38),
            when,
            12,
            italic=True,
            color=when_color,
            align=PP_ALIGN.CENTER,
        )
        _textbox(slide, (x - col_w / 2, 2.48, col_w, 0.40), num, 20, bold=True, align=PP_ALIGN.CENTER)
        _textbox(slide, (x - col_w / 2, 2.90, col_w, 0.42), title, 16, bold=True, align=PP_ALIGN.CENTER)
        _textbox(slide, (x - col_w / 2, 3.32, col_w, 0.55), sub, 14, color=GRAY, align=PP_ALIGN.CENTER)

    caption(slide, "Тест блока А — после Л4, не сегодня. Вторая половина Л4 уже открывает блок Б.")
    return slide


def set_table_fonts(table, header_rows=1, size=12):
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _run(run, size, bold=(i < header_rows or j == 0))
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALE if i < header_rows else WHITE


def add_table(slide, rows, left, top, width, height, col_widths, header_rows=1, size=12):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    set_table_fonts(table, header_rows=header_rows, size=size)
    return table


def build():
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    make_figures()
    prs = Presentation(str(SRC))
    delete_all_slides(prs)

    title_slide(
        prs,
        "Вычислительная цветовая фотография",
        "Лекция 3. Линейная модель в деталях",
        "IITP'26  ·  ИППИ РАН  ·  осень 2026",
    )

    slide_block_a_timeline(prs)

    section(
        prs,
        "Формула и её границы",
        photo=KHARK / "kh_1_003s.jpg",
        note="Из чего собрана линейная модель и где она кончается.",
    )

    s = titled(prs, "Что уже есть")
    pic(s, FIG / "eq_sensor.png", 1.05, 1.22, width=7.9)
    set_body(
        s,
        [
            "Свет дошёл как спектр-стимул F; в конце его взвешивает канал χᵢ.",
            "У колбочки, аномала и камеры разные χ, форма интеграла одна.",
            "На Л2 для матовой поверхности писали F = I · R. Сегодня этот случай раскрываем.",
        ],
        top=2.55,
        height=2.15,
        size=16,
    )

    s = titled(prs, "Полная запись")
    fit_pic(s, FIG / "eq_full_stack.png", (0.35, 1.08, 9.30, 1.48))
    set_body(
        s,
        [
            "Источники: спектр S_j и диаграмма B_j; потоки складываются.",
            "Поверхность: несколько R_k и геометрия f, в том числе (n · n_i)_+.",
            "Сенсор в конце: у c_i нужен индекс n_r.",
        ],
        top=2.62,
        width=BODY_LEFT_W,
        height=2.20,
        size=15,
    )
    fit_pic(s, FIG / "l12_brdf.png", (5.27, 2.58, 4.35, 2.20))
    caption(s, "Схема с Л1–Л2. Овал Φ — окраска на пятне; стимул к глазу — F.")

    s = titled(prs, "Что в модели линейно")
    hide_body(s)
    fit_pic(s, FIG / "eq_full_F.png", (0.45, 1.10, 9.10, 1.18))
    add_table(
        s,
        [
            ("Член", "Линейно по", "Физика"),
            ("S_j, B_j и F", "сложение потоков", "два источника → сумма стимулов (Грассман)"),
            ("R_k", "смесь компонент", "сколько угодно слагаемых: пигменты, лак, блик"),
            ("χᵢ", "пока χ фиксирован", "адаптация и мезопическое зрение выключают эту строку"),
            ("∫ dλ, ∫ dω", "суперпозиция", "интеграл — линейный оператор"),
        ],
        0.45,
        2.38,
        9.1,
        2.40,
        [2.0, 2.3, 4.8],
    )

    four_pic(
        prs,
        "Где модель ломается",
        [
            (PHOTO / "surf_cd.jpg", "дифракция, интерференция"),
            (PHOTO / "fluor.jpg", "флуоресценция: λ_in ≠ λ_out"),
            (PHOTO / "night.jpg", "адаптация и мезопическое зрение"),
            (PHOTO / "skin.jpg", "полупрозрачные среды, BSSRDF"),
        ],
    )

    section(
        prs,
        "Источники",
        photo=KHARK / "kh_1_014s.jpg",
        note="У источника два входа: геометрия B_j и спектр S_j(λ).",
    )

    s = titled(prs, "Геометрия источника")
    hide_body(s)
    add_table(
        s,
        [
            ("Тип", "Идеализация", "Что входит в S_j B_j"),
            ("Точечный / направленный", "δ(ω−ω₀)", "одно направление, жёсткие тени"),
            ("Линейный", "отрезок / цилиндр", "протяжённый блик, мягкая тень вдоль оси"),
            ("Площадной", "диск, окно, небо", "интеграл по Ω₊, полутени"),
            ("Удалённый (солнце)", "параллельные лучи", "как точка на бесконечности"),
            ("Рассеянный / амбиент", "почти изотропный", "мало направления, много заливки"),
        ],
        0.40,
        1.22,
        9.2,
        3.55,
        [2.6, 2.4, 4.2],
    )

    s = titled(prs, "Типовые источники")
    hide_body(s)
    fit_pic(s, FIG / "typical_sources.png", (0.40, 1.14, 9.20, 3.70))
    caption(s, "CIE 15:2018. Натуральные планковские vs искусственные: FL, LED-B, LED-V, RGB.")

    s = titled(prs, "Стандарты CIE")
    hide_body(s)
    add_table(
        s,
        [
            ("Семья", "Имя", "Смысл"),
            ("планк", "A", "2856 K, накал"),
            ("планк", "E", "S(λ) = const, теория"),
            ("серия D", "D50", "≈ 5003 K, печать, ICC"),
            ("серия D", "D55, D65, D75", "дневной; D65 — колориметрия и sRGB"),
            ("серия D", "ID50, ID65", "дневной за стеклом"),
            ("FL", "FL2, FL7, FL11", "трубки; TL84 — три пика"),
            ("LED", "LED-B1 … B5", "синий кристалл + люминофор"),
            ("LED", "LED-BH1", "плюс красный диод"),
            ("LED", "LED-RGB1", "три узких пика"),
            ("LED", "LED-V1, V2", "фиолетовая накачка, широкополосные"),
        ],
        0.40,
        1.16,
        9.2,
        3.55,
        [1.7, 2.5, 5.0],
        size=11,
    )
    caption(s, "CIE 15:2018. Серия D — одна семья. Конкретная лампа равна стандарту, только если совпал спектр.")

    s = titled(prs, "Источник в формуле сцены")
    pic(s, FIG / "eq_diffuse.png", 0.90, 1.30, width=4.8)
    set_body(
        s,
        [
            "Свёртка S с окраской Φ, когда геометрия уже спрятана в Φ.",
            "В общем случае — интеграл полной записи: по каким направлениям и с какими S_j, B_j копить СДФОС.",
        ],
        top=2.55,
        width=4.80,
        height=1.85,
        size=16,
    )
    cover_pic(s, PHOTO / "surf_apple.jpg", PIC_RIGHT)

    s = titled(prs, "Поверхности и СДФОС")
    set_body(
        s,
        [
            "СДФОС = BRDF: R(n_i, n_r, λ). Единица — 1/ср.",
            "Ламберт, как у Максимова: R = Φ(λ)/π, угла в спектре почти нет.",
            "Блик, металл, бархат, хамелеон — нужна функция двух лучей.",
            "f — площадка и телесный угол; R — материал. Компонент R_k сразу несколько.",
        ],
        width=5.85,
        size=16,
    )
    cover_pic(s, KHARK / "kh_02_01s.jpg", (6.70, 1.16, 3.15, 3.85))
    caption(s, "Фотоархив А. А. Харкевича, ИППИ РАН")

    content(
        prs,
        "Ограничения на R",
        [
            "Неотрицательность. Взаимность Гельмгольца.",
            "Энергия: интеграл по полусфере с косинусом ≤ 1.",
            "Без флуоресценции одна и та же λ на входе и на выходе.",
        ],
    )

    four_pic(
        prs,
        "Частные случаи",
        [
            (PHOTO / "surf_paper.jpg", "бумага: ламберт, R = Φ/π"),
            (PHOTO / "surf_chrome.jpg", "хром и металлы: Френель; Au, Cu красят блик"),
            (PHOTO / "surf_apple.jpg", "яблоко: тело + зеркальный лепесток"),
            (PHOTO / "surf_velvet.jpg", "бархат: угол без зеркального лепестка"),
        ],
        caption_txt="Хамелеон: спектр едет с углом. Cook–Torrance — сумма нескольких R_k.",
    )

    text_pic(
        prs,
        "Когда СДФОС мало",
        [
            "Флуоресценция — биспектральная R(λ_in, λ_out).",
            "BSSRDF: кожа, мрамор, воск — вход и выход в разных точках.",
            "Объём, интерференция, дифракция (крыло, CD).",
        ],
        PHOTO / "surf_cd.jpg",
        cover=True,
    )

    section(
        prs,
        "Сенсор: глаз",
        photo=KHARK / "kh_02_08s.jpg",
        note="Приёмник χᵢ: оптика до рецептора, колбочки, сетчатка, аномалии, средний CIE и индивид.",
    )

    s = titled(prs, "Сетчатка, как она выглядит")
    set_body(
        s,
        [
            "Фундус: диск нерва, макула, сосуды.",
            "Свет приходит со стороны стекловидного: сначала ганглии и аксоны.",
            "Рецепторы глубже, у пигментного эпителия; наружные сегменты смотрят от света.",
            "Фотон проходит сосуды и проводку, потом попадает в сегмент.",
            "В ямке передний слой тоньше; макула желтит до колбочки.",
        ],
        width=BODY_LEFT_W,
        size=14,
    )
    cover_pic(s, PHOTO / "eye_fundus.jpg", (5.27, 1.22, 4.35, 1.68))
    fit_pic(s, CROP / "eye_histo.png", (5.27, 2.98, 4.35, 1.68))
    caption(s, "Сверху: фундус, NEI, PD. Снизу: слои, OpenStax CC BY.")

    text_pic(
        prs,
        "Оптика до рецептора",
        [
            "Роговица, зрачок, хрусталик, макула — фильтр до колбочки.",
            "Меняется F, дошедшее до пигмента.",
            "Поле 2° — ямка. Поэтому Райт и Гилд брали узкое поле.",
        ],
        FIG / "l12_eye.png",
        caption_txt="Та же схема глаза, что на Л1–Л2.",
    )

    text_pic(
        prs,
        "Фоторецепторы",
        [
            "Палочки: родопсин, скотопическое зрение; один тип — нет трёхмерных равенств.",
            "Колбочки L, M, S: три пигмента, сильное перекрытие L и M.",
            "Мезопическое зрение — смесь. Для равенств палочки выключают яркостью и полем.",
        ],
        FIG / "l12_lms.png",
        caption_txt="LMS с Л1–Л2.",
    )

    text_pic(
        prs,
        "Цепочка в сетчатке",
        [
            "Колбочка: на свету гиперполяризация, меньше глутамата. Здесь живут равенства Максвелла.",
            "Горизонтальные вычитают фон соседей. Биполяры ON/OFF кодируют перепад.",
            "Амакриновые — время и движение. Ганглии — спайки: L−M, S−(L+M), яркость.",
        ],
        CROP / "eye_circuit.png",
        caption_txt="OpenStax. Оппонентность — мост к Л4.",
    )

    content(
        prs,
        "ipRGC / меланопсин",
        [
            "Нужны зрачку и циркадности.",
            "Равенство Максвелла ими не ставят.",
            "Размерность равенств задаёт число независимых линейных функционалов в поле Максвелла.",
        ],
    )

    s = titled(prs, "Аномалии: другие χᵢ")
    hide_body(s)
    add_table(
        s,
        [
            ("", "Нет пигмента", "Пигмент сдвинут"),
            ("L", "протанопия", "протаномалия"),
            ("M", "дейтеранопия", "дейтераномалия"),
            ("S", "тританопия", "тританомалия"),
        ],
        0.40,
        1.22,
        5.05,
        2.35,
        [0.85, 2.10, 2.10],
    )
    set_body(
        s,
        [
            "Другие χᵢ ⇒ другая линейная модель и другие метамеры.",
            "L и M — X-хромосома. Диагностика аномалоскопом — на Л4.",
        ],
        left=0.40,
        top=3.70,
        width=5.05,
        height=1.05,
        size=15,
    )
    cover_pic(s, PHOTO / "eye_mosaic.jpg", (5.65, 1.22, 3.95, 3.45))
    caption(s, "Мозаика колбочек, Fairchild / Commons.")

    content(
        prs,
        "Три разных наблюдателя",
        [
            "Живой глаз и аномал — разные χ на одном F.",
            "CIE 1931 2° и CIE 1964 10° — средний χ, не чей-то глаз.",
            "Нюберг писал Райту: усреднённый наблюдатель никому не равен.",
            "Камера — ещё один набор. Разброс нормальных — модель Асано.",
        ],
    )

    s = titled(prs, "Модель Асано")
    set_body(
        s,
        [
            "Линза режет коротковолновую часть, в основном χ_S.",
            "Макула: яма около 460 нм; поле 2° плотнее, чем 10°.",
            "Плотность пигмента: кривая χ шире, L и M сильнее перекрываются.",
            "Сдвиг λ_max: номограмма едет по оси. Это разброс нормальных.",
        ],
        width=BODY_LEFT_W,
        size=15,
    )
    fit_pic(s, FIG / "asano_chi.png", PIC_RIGHT)
    caption(s, "Asano, Fairchild, Blondé, PLOS ONE 2016. Надстройка над CIEPO06.")

    section(
        prs,
        "Из линейной модели",
        photo=KHARK / "kh_02_27s.jpg",
        note=[
            "Райт и Гилд.",
            "Выпуклый конус.",
            "Цветовое тело.",
            "Пересечения и метамерия.",
            "Ступенчатые окраски.",
        ],
    )

    section(
        prs,
        "1. Райт и Гилд",
        photo=KHARK / "kh_1_001s.jpg",
        note="Почему два опыта складывают в одного наблюдателя. Канг, §2.8–2.9.",
    )

    text_pic(
        prs,
        "Два опыта, разные первичные",
        [
            "Guild, 1931: 7 человек, 2°, фильтры.",
            "Wright: 10 человек, 2°, спектральные первичные ≈ 650, 530, 460 нм.",
            "CIE 1931: смесь 7+10, общие 700 / 546.1 / 435.8 нм.",
        ],
        FIG / "l12_match.png",
        caption_txt="Схема равенства с Л1–Л2. Ohta, Note 3.1; Канг, гл. 2.",
    )

    s = titled(prs, "Зачем их можно совмещать")
    fit_pic(s, FIG / "eq_kang.png", (0.40, 1.10, 9.20, 2.05))
    set_body(
        s,
        [
            "Два набора CMF связаны одной 3×3. Новой физики нет: Грассман и замена базиса.",
            "Поэтому Райт и Гилд пересчитываются в общие CIE RGB. Оба — нормальный трихромат, 2°, ямка.",
            "Если Грассман не держит, матрица (2.31) уже не про того же наблюдателя.",
        ],
        top=3.22,
        height=1.55,
        size=15,
    )
    caption(s, "Канг, ур. 2.29–2.31. Φ_j здесь — спектры первичных, не окраска.")

    s = content(
        prs,
        "Как из хроматичности делают CMF",
        [
            "Локус в своих первичных → матрица (2.31) в CIE RGB.",
            "Масштаб по V(λ), среднее Guild + Wright.",
            "Среднее векторов равенства — снова вектор равенства. Это тот средний, про которого Нюберг писал Райту.",
        ],
    )
    pause(s, "Пауза. Можно ли так же усреднить протанопа и нормального?")

    section(
        prs,
        "2. Выпуклый цветовой конус",
        photo=KHARK / "kh_1_023s.jpg",
        note="Неотрицательные спектры. Линейный образ — выпуклый конус.",
    )

    content(
        prs,
        "Что называют конусом",
        [
            "Спектральные: F(λ) = E₀ δ(λ−λ₀). Луч из нуля при смене E₀.",
            "Локус замыкают пурпурной хордой (смесь крайних λ).",
            "Конус — цвета всех F ≥ 0, не только монохроматов. Выпуклая оболочка локуса.",
        ],
    )

    two_pic(
        prs,
        "Конус, локус, треугольник",
        FIG / "l12_cone.png",
        FIG / "l12_xy.png",
        caption_txt="Как на Л1–Л2: конус в XYZ и хроматический локус.",
        b1=(0.35, 1.14, 5.25, 3.70),
        b2=(5.80, 1.22, 3.80, 3.50),
    )

    pic_slide(
        prs,
        "Глаз и камера — разный локус",
        CROP / "paper_fig1.png",
        caption_txt="Kroshnin, Vasilev, Ershov et al., JOSA A 39:452 (2022), рис. 1. Слева CIE, справа Nikon D90.",
        box=(0.35, 1.15, 9.30, 3.55),
    )

    s = content(
        prs,
        "Почему выпуклый — из Грассмана",
        [
            "{F ≥ 0} — выпуклый конус спектров-стимулов.",
            "C линейно. Образ выпуклого конуса — выпуклый конус в R³.",
            "У камеры локус может быть невыпуклым — треугольник всё равно выпуклый.",
        ],
    )
    pause(s, "Пауза. Отрицательный спектр — цвет? Почему нуль — вершина?")

    section(
        prs,
        "3. Цветовое тело",
        photo=KHARK / "kh_02_21s.jpg",
        note="Несамосветящиеся окраски при данном освещении.",
    )

    s = titled(prs, "Определение")
    pic(s, FIG / "eq_diffuse.png", 0.90, 1.22, width=4.8)
    set_body(
        s,
        [
            "0 ≤ Φ(λ) ≤ 1, F = S Φ. Тело при данном S — все цвета таких окрасок.",
            "Лежит внутри конуса и ограничено: 0 ≤ Xᵢ ≤ Xᵢ⁰, где X⁰ = C(S).",
            "Форма зависит от S и χᵢ. У конуса ламп нет потолка Φ ≤ 1.",
        ],
        top=2.55,
        height=2.00,
        size=16,
    )

    pic_slide(
        prs,
        "Тела зависят от света и от χᵢ",
        CROP / "maximov_solids.png",
        caption_txt="Максимов, рис. 27: тела и метамерные области при разных освещениях.",
        box=(0.45, 1.12, 9.10, 3.55),
    )

    s = titled(prs, "Центральная симметрия")
    pic(s, FIG / "eq_complement.png", 0.7, 1.18, width=8.5)
    set_body(
        s,
        [
            "Φ* = 1−Φ. Середина отрезка — 50% серый, ½ X⁰.",
            "У конуса ламп такой пары нет: S−F может выйти в отрицательный спектр.",
        ],
        top=2.55,
        height=1.6,
        size=16,
    )
    caption(s, "Максимов: дополнительные окраски.")

    text_pic(
        prs,
        "Двойной конус — грубая форма",
        [
            "Тело внутри конуса из 0 и конуса из X⁰.",
            "Это оценка. Точная граница — ступеньки.",
        ],
        CROP / "paper_fig1_cie.png",
    )

    section(
        prs,
        "4. Пересечения и метамерия",
        photo=KHARK / "kh_1_007s.jpg",
        note="Когда по графикам сразу видно, что метамерии нет.",
    )

    s = titled(prs, "Постановка")
    pic(s, FIG / "eq_delta.png", 0.55, 1.22, width=4.7)
    set_body(
        s,
        ["Δ = F − Ψ ортогональна всем χᵢ. Когда по графикам сразу видно, что метамерии нет?"],
        top=2.55,
        width=4.8,
        height=1.5,
        size=16,
    )
    fit_pic(s, FIG / "l12_lettuce.png", (5.50, 1.22, 4.10, 3.45))
    caption(s, "Спектры объектов и LMS — слайд с Л1–Л2 (Brown / lettuce).")

    two_pic(
        prs,
        "Метамеры на фото",
        FIG / "l12_lettuce.png",
        FIG / "l12_lettuce_ink.png",
        caption_txt="Салат vs зелёные чернила: разные F, близкий отклик. Л1–Л2.",
        b1=(0.45, 1.18, 4.55, 3.55),
        b2=(5.20, 1.18, 4.40, 3.55),
    )

    content(
        prs,
        "Ноль и одно пересечение",
        [
            "Δ не меняет знак: все χᵢ видят больше — метамерии нет.",
            "Одно пересечение может обнулить один канал, не два независимых.",
            "Меньше двух пересечений ⇒ не метамеры. Достаточное условие.",
        ],
    )

    s = titled(prs, "Число пересечений")
    pic(s, FIG / "crossings.png", 0.40, 1.18, width=9.2)
    set_body(
        s,
        ["Меньше двух пересечений ⇒ не метамеры. Глазу обычно нужно три."],
        top=3.85,
        height=0.85,
        size=16,
        bullets=False,
    )
    caption(s, "Stiles & Wyszecki, JOSA 1968. Стиль кривых — как в JOSA A 39:452.")

    s = pic_slide(
        prs,
        "Свет меняет тело и пересечения",
        CROP / "maximov_metamers.png",
        caption_txt="Максимов, рис. 45: спектры освещения и перекос цветового тела.",
        box=(0.45, 1.12, 9.10, 3.40),
    )
    pause(s, "Пауза. Три пересечения обязаны быть метамерами? Нет.")

    section(
        prs,
        "5. Ступенчатые окраски",
        photo=KHARK / "kh_1_026s.jpg",
        note="Оболочка тела и любая точка внутри.",
    )

    s = titled(prs, "Что такое ступенька")
    pic(s, FIG / "steps.png", 0.40, 1.18, width=9.2)
    set_body(
        s,
        ["Φ ∈ {0, 1}, m скачков. На оболочке глаза m = 0, 1, 2."],
        top=3.85,
        height=0.85,
        size=16,
    )

    pic_slide(
        prs,
        "Утверждение 1. Оболочка тела",
        CROP / "maximov_steps_solid.png",
        caption_txt="Максимов, рис. 18: одноступенчатые vs поверхность тела; врезки — Φ ступеньки.",
        box=(0.40, 1.12, 9.20, 3.55),
    )

    s = content(
        prs,
        "Почему на границе нет значений между 0 и 1",
        [
            "Цвет на поверхности, но 0 < Φ < 1 на малом интервале.",
            "Сжать в 0 и в 1: исходная — смесь концов, внутри отрезка.",
            "Значит на поверхности Φ ∈ {0, 1}. Для глаза m = 0, 1, 2.",
        ],
    )
    pause(s, "Пауза. Серая Φ ≡ 1/2 на поверхности? Нет: это центр.")

    content(
        prs,
        "Как пробегать оболочку",
        [
            "Полоса 0→1→0 или стоп 1→0→1, двигаем λ₁ < λ₂.",
            "Дополнение 1−Φ — вторая створка.",
            "Схлопнуть полосу — ребро. Всё 0 / всё 1 — вершины.",
        ],
    )

    s = titled(prs, "Утверждение 2. Смесь ступенек")
    pic(s, FIG / "eq_mix.png", 0.8, 1.25, width=8.3)
    set_body(
        s,
        [
            "Любая точка тела — выпуклая комбинация ступенек.",
            "Сложение — усреднение окрасок. Внутри точки метамерны.",
        ],
        top=2.70,
        height=1.50,
        size=16,
    )
    caption(s, "В статье 2022 это banded model: спектральная модель, замкнутая и по сложению, и по умножению.")

    s = titled(prs, "Явный способ: кусочно-постоянная Φ")
    pic(s, FIG / "eq_piecewise.png", 1.0, 1.22, width=8.0)
    set_body(
        s,
        [
            "Каждый 1_{I_k} — прямоугольная ступенька.",
            "Так приближаем любую Φ; внутри точки метамерны.",
        ],
        top=2.55,
        height=1.60,
        size=16,
    )

    s = titled(prs, "Что не смешивать")
    hide_body(s)
    add_table(
        s,
        [
            ("", "Ступенька одной Φ", "Пересечения двух спектров"),
            ("Вопрос", "где точка тела", "метамеры или нет"),
            ("Число", "m < 3 на оболочке глаза", "обычно ≥ 3, чтобы метамерия была возможна"),
            ("Сложение", "смесь даёт внутренность", "разность Δ в ядре CMF"),
        ],
        0.45,
        1.35,
        9.1,
        2.9,
        [1.7, 3.7, 3.7],
    )

    section(
        prs,
        "Что мы за сегодня поняли?",
        photo=KHARK / "kh_02_24s.jpg",
        note="Формула по ходу света и пять следствий линейности.",
    )

    content(
        prs,
        "Что сегодня зафиксировали",
        [
            "S_j B_j, R_k, f, F, χᵢ — в таком порядке. Компонент R_k может быть больше двух.",
            "Источники: геометрия отдельно, спектр отдельно. Планк vs FL/LED, CIE 15:2018.",
            "Окраска Φ; общее — СДФОС двух лучей. Глаз: колбочки, Асано правит форму χ.",
            "Райт+Гилд: Канг (2.31). Конус выпуклый. Тело: Φ ↔ 1−Φ. Оболочка — ступеньки.",
        ],
    )
    content(
        prs,
        "Дальше",
        [
            "Л4: пороги (Раутиан), адаптация, CAM — где линейное равенство кончается.",
            "Тест блока А — после Л4. CST и камера — блок Б.",
        ],
    )
    content(
        prs,
        "К зачёту",
        [
            "Карточки А2.5–7, 9 — Грассман, измерение, Максвелл, CMF.",
            "А3.10–12 — параметризация, конус и тело, лемма о метамерах.",
            "А2.8 и Раутиан — после Л4. А3.13 (Планк/CCT) — не на эту пару.",
        ],
    )

    try:
        prs.save(str(OUT))
        saved = OUT
    except PermissionError:
        saved = OUT.with_name(OUT.stem + "_new.pptx")
        prs.save(str(saved))
    return prs, saved


if __name__ == "__main__":
    _, saved = build()
    print("wrote", saved.name)
